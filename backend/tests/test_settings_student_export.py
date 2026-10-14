from io import BytesIO

import openpyxl

from models import Student
from tests.conftest import CHEM_ADMIN_EMAIL
from tests.test_settings_student_import import THAI_HEADERS, upload


URL = "/api/admin/students/export"


def add_student(db_session, student_id, major, first_name="สมชาย", email="a@cmu.ac.th"):
    db_session.add(
        Student(
            student_id=student_id,
            first_name=first_name,
            last_name="ใจดี",
            email=email,
            faculty="คณะวิทยาศาสตร์",
            major=major,
        )
    )
    db_session.commit()


def rows_of(response):
    sheet = openpyxl.load_workbook(BytesIO(response.content)).active
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def test_export_uses_template_headers_and_all_students(client, db_session):
    add_student(db_session, "680510002", "วิทยาการคอมพิวเตอร์")
    add_student(db_session, "680510001", "เคมี")

    response = client.get(URL)

    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    rows = rows_of(response)
    assert rows[0] == THAI_HEADERS
    assert rows[1:] == [
        ["680510001", "สมชาย", "ใจดี", "a@cmu.ac.th", "เคมี"],
        ["680510002", "สมชาย", "ใจดี", "a@cmu.ac.th", "วิทยาการคอมพิวเตอร์"],
    ]


def test_exported_file_imports_back(client, db_session):
    add_student(db_session, "680510001", "เคมี")
    exported = client.get(URL).content

    db_session.query(Student).delete()
    db_session.commit()

    body = upload(client, exported).json()
    assert (body["total"], body["success"], body["failed"]) == (1, 1, 0)
    db_session.expire_all()
    assert db_session.get(Student, "680510001").first_name == "สมชาย"


def test_values_starting_with_equals_stay_text(client, db_session):
    add_student(db_session, "680510001", "เคมี", first_name="=HYPERLINK(\"x\")")

    sheet = openpyxl.load_workbook(BytesIO(client.get(URL).content)).active

    assert sheet["B2"].data_type == "s"
    assert sheet["B2"].value == "=HYPERLINK(\"x\")"


def test_branch_admin_exports_only_own_major(client, db_session):
    add_student(db_session, "680510001", "เคมี")
    add_student(db_session, "680510002", "วิทยาการคอมพิวเตอร์")

    rows = rows_of(client.get(URL, headers={"X-Admin-Email": CHEM_ADMIN_EMAIL}))

    assert [row[0] for row in rows[1:]] == ["680510001"]


def test_export_requires_admin(client):
    assert client.get(URL, headers={"X-Admin-Email": ""}).status_code == 401
